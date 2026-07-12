"""Generate a downloadable results report (Word .docx + Markdown) from the
evaluated checkpoints — the benchmark tables the project produces, captured in
one shareable document.

Computes, for the bicubic floor and each checkpoint: overall PSNR/SSIM, the
per-terrain breakdown, and the per-terrain delta of the terrain-aware model
(passed last) minus each baseline. Reads dataset composition from the split
manifest. Reuses the real evaluation code (evaluation/_eval_common.py), so the
numbers match `eval_psnr_ssim.py` / `eval_per_terrain.py` exactly.

Writes <out>.md always (no dependencies) and <out>.docx if python-docx is
installed (falls back to Markdown-only with a note otherwise).

    python evaluation/make_results_report.py \
        --test-csv data/dataset/test.csv \
        --split-manifest data/dataset/dataset_manifest_split.csv \
        --checkpoints checkpoints/srcnn/best.pth checkpoints/srgan/best.pth \
                      checkpoints/swinir/best.pth checkpoints/terrasr/best.pth \
        --out data/dataset/results_report
"""
import argparse
import datetime as dt
import sys
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from evaluation._eval_common import (load_model_from_checkpoint,  # noqa: E402
                                       make_test_loader, run_sr_over_loader)
from training.train_utils import get_device  # noqa: E402


def compute_all(test_csv, terrain_config, checkpoints, with_bicubic, bicubic_scale):
    device = get_device()
    loader = make_test_loader(test_csv, terrain_config)
    n = len(loader.dataset)
    results = {}   # label -> per-patch DataFrame
    if with_bicubic:
        results["bicubic"] = run_sr_over_loader(loader, device, bicubic_scale=bicubic_scale)
    for ckpt in checkpoints:
        model, name, is_terrain = load_model_from_checkpoint(ckpt, device)
        results[name] = run_sr_over_loader(loader, device, model=model, is_terrain=is_terrain)
    return results, n, str(device)


def overall_rows(results):
    return [(label, df["psnr"].mean(), df["ssim"].mean()) for label, df in results.items()]


def per_terrain(df):
    return df.groupby("terrain").agg(n=("psnr", "size"), psnr=("psnr", "mean"),
                                      ssim=("ssim", "mean")).sort_index()


def dataset_composition(split_manifest):
    if split_manifest is None or not Path(split_manifest).exists():
        return None
    df = pd.read_csv(split_manifest) if str(split_manifest).endswith(".csv") \
        else pd.read_parquet(split_manifest)
    comp = {
        "total": len(df),
        "scenes": int(df["source_scene"].nunique()) if "source_scene" in df else None,
        "terrain_counts": df["terrain_label"].value_counts().to_dict() if "terrain_label" in df else {},
        "split_counts": df["split"].value_counts().to_dict() if "split" in df else {},
        "pan": df["pseudo_pan"].value_counts().to_dict() if "pseudo_pan" in df else {},
    }
    return comp


# ---------------- Markdown ----------------

def md_table(headers, rows):
    out = ["| " + " | ".join(headers) + " |",
           "| " + " | ".join("---" for _ in headers) + " |"]
    for r in rows:
        out.append("| " + " | ".join(str(c) for c in r) + " |")
    return "\n".join(out)


def build_markdown(results, n_test, device, comp, title):
    L = [f"# {title}", "",
         f"_Generated {dt.datetime.now():%Y-%m-%d %H:%M}  ·  test patches: {n_test}  ·  device: {device}_",
         ""]
    if comp:
        L += ["## Dataset composition", "",
              f"- Total patches: **{comp['total']}**" + (f"  ·  scenes: {comp['scenes']}" if comp['scenes'] else ""),
              md_table(["Terrain", "Patches"], list(comp["terrain_counts"].items())), ""]
        if comp["split_counts"]:
            L += ["Split: " + ", ".join(f"{k} {v}" for k, v in comp["split_counts"].items()), ""]
        if comp["pan"]:
            L += ["PAN provenance: " + ", ".join(
                f"{'pseudo-PAN' if str(k).lower()=='true' else 'true PAN'} {v}"
                for k, v in comp["pan"].items()), ""]

    L += ["## Overall PSNR / SSIM", "",
          md_table(["Model", "PSNR (dB)", "SSIM"],
                   [(m, f"{p:.3f}", f"{s:.4f}") for m, p, s in overall_rows(results)]), ""]

    L += ["## Per-terrain breakdown", ""]
    for label, df in results.items():
        pt = per_terrain(df)
        L += [f"**{label}**", "",
              md_table(["Terrain", "n", "PSNR (dB)", "SSIM"],
                       [(idx, int(r.n), f"{r.psnr:.3f}", f"{r.ssim:.3f}") for idx, r in pt.iterrows()]), ""]

    labels = list(results)
    if len(labels) >= 2:
        ref = labels[-1]
        ref_pt = per_terrain(results[ref])
        L += [f"## Per-terrain delta ({ref} − baseline)", ""]
        for other in labels[:-1]:
            j = per_terrain(results[other]).join(ref_pt, lsuffix="_b", rsuffix="_r")
            j["dPSNR"] = j["psnr_r"] - j["psnr_b"]
            j["dSSIM"] = j["ssim_r"] - j["ssim_b"]
            L += [f"**{ref} − {other}**", "",
                  md_table(["Terrain", "Δ PSNR (dB)", "Δ SSIM"],
                           [(idx, f"{r.dPSNR:+.3f}", f"{r.dSSIM:+.3f}") for idx, r in j.iterrows()]), ""]

    L += ["---",
          "_Note: numbers reflect the dataset and training configuration used for this run. "
          "Small-scale / few-epoch runs are for pipeline verification and are not "
          "representative benchmark results._"]
    return "\n".join(L)


# ---------------- DOCX ----------------

def build_docx(results, n_test, device, comp, title, out_docx):
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
    except ImportError:
        return False

    ACCENT = RGBColor(0x1F, 0x6B, 0x66)
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    for lvl, sz in [("Heading 1", 15), ("Heading 2", 12)]:
        st = doc.styles[lvl]; st.font.color.rgb = ACCENT; st.font.size = Pt(sz); st.font.bold = True

    def shade(cell, hexfill, white=False):
        tcpr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), hexfill); tcpr.append(sh)

    def table(headers, rows):
        t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i]; c.text = ""; run = c.paragraphs[0].add_run(str(h))
            run.bold = True; run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); shade(c, "1F6B66")
        for r_i, row in enumerate(rows):
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = str(v)
                if r_i % 2 == 1:
                    shade(cells[i], "EDF2F1")
        doc.add_paragraph()

    h = doc.add_heading(title, level=0)
    doc.add_paragraph(f"Generated {dt.datetime.now():%Y-%m-%d %H:%M}  ·  test patches: {n_test}  ·  device: {device}")

    if comp:
        doc.add_heading("Dataset composition", level=1)
        doc.add_paragraph(f"Total patches: {comp['total']}"
                          + (f"   ·   scenes: {comp['scenes']}" if comp['scenes'] else ""))
        table(["Terrain", "Patches"], list(comp["terrain_counts"].items()))

    doc.add_heading("Overall PSNR / SSIM", level=1)
    table(["Model", "PSNR (dB)", "SSIM"],
          [(m, f"{p:.3f}", f"{s:.4f}") for m, p, s in overall_rows(results)])

    doc.add_heading("Per-terrain breakdown", level=1)
    for label, df in results.items():
        doc.add_heading(label, level=2)
        pt = per_terrain(df)
        table(["Terrain", "n", "PSNR (dB)", "SSIM"],
              [(idx, int(r.n), f"{r.psnr:.3f}", f"{r.ssim:.3f}") for idx, r in pt.iterrows()])

    labels = list(results)
    if len(labels) >= 2:
        ref = labels[-1]; ref_pt = per_terrain(results[ref])
        doc.add_heading(f"Per-terrain delta ({ref} minus baseline)", level=1)
        for other in labels[:-1]:
            j = per_terrain(results[other]).join(ref_pt, lsuffix="_b", rsuffix="_r")
            j["dPSNR"] = j["psnr_r"] - j["psnr_b"]; j["dSSIM"] = j["ssim_r"] - j["ssim_b"]
            doc.add_heading(f"{ref} − {other}", level=2)
            table(["Terrain", "Δ PSNR (dB)", "Δ SSIM"],
                  [(idx, f"{r.dPSNR:+.3f}", f"{r.dSSIM:+.3f}") for idx, r in j.iterrows()])

    note = doc.add_paragraph()
    run = note.add_run("Note: numbers reflect the dataset and training configuration used for "
                       "this run. Small-scale / few-epoch runs are for pipeline verification "
                       "and are not representative benchmark results.")
    run.italic = True
    doc.save(out_docx)
    return True


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--test-csv", required=True, type=Path)
    ap.add_argument("--checkpoints", nargs="+", required=True, type=Path)
    ap.add_argument("--split-manifest", type=Path, default=None,
                     help="dataset_manifest_split.csv/.parquet for the composition section")
    ap.add_argument("--terrain-config", default="configs/terrain_classes.yaml")
    ap.add_argument("--with-bicubic", action="store_true", default=True)
    ap.add_argument("--no-bicubic", dest="with_bicubic", action="store_false")
    ap.add_argument("--bicubic-scale", type=int, default=2)
    ap.add_argument("--title", default="TerraSR — Results Report")
    ap.add_argument("--out", required=True, type=Path,
                     help="output path WITHOUT extension (writes .md and .docx)")
    args = ap.parse_args()

    results, n_test, device = compute_all(
        args.test_csv, args.terrain_config, args.checkpoints, args.with_bicubic, args.bicubic_scale)
    comp = dataset_composition(args.split_manifest)

    args.out.parent.mkdir(parents=True, exist_ok=True)
    md_path = args.out.with_suffix(".md")
    md_path.write_text(build_markdown(results, n_test, device, comp, args.title), encoding="utf-8")
    print(f"wrote {md_path}")

    docx_path = args.out.with_suffix(".docx")
    if build_docx(results, n_test, device, comp, args.title, str(docx_path)):
        print(f"wrote {docx_path}")
    else:
        print("python-docx not installed — wrote Markdown only "
              "(pip install python-docx for the .docx report)")


if __name__ == "__main__":
    main()
